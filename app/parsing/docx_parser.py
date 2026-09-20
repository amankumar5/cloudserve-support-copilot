import os
import uuid
import docx
from typing import List
from app.models.document import Document, Page, DocumentElement, ElementType
from app.parsing.base import BaseDocumentParser
from app.parsing.vision_analyzer import vision_analyzer
from app.core.config import settings
from app.core.logging import logger


class DOCXParser(BaseDocumentParser):
    """
    DOCX Document Parser extracting text headings, paragraphs, tables, and inline images.
    """

    def parse(self, file_path: str, file_id: str, file_name: str, gdrive_url: str = "") -> Document:
        logger.info(f"Parsing DOCX document: {file_name}")
        document_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, file_id or file_name))

        doc_obj = docx.Document(file_path)
        elements: List[DocumentElement] = []
        current_section = "Main Section"
        elem_counter = 0

        # Process paragraphs
        for para in doc_obj.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            elem_counter += 1
            if para.style.name.startswith("Heading"):
                current_section = text
                elem_type = ElementType.HEADING
            else:
                elem_type = ElementType.TEXT

            elements.append(DocumentElement(
                element_id=f"{document_id}_docx_elem{elem_counter}",
                element_type=elem_type,
                page_number=1,
                section_title=current_section,
                content_text=text
            ))

        # Process tables
        for tbl_idx, table in enumerate(doc_obj.tables):
            elem_counter += 1
            raw_rows = []
            for row in table.rows:
                row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                raw_rows.append(row_data)

            if not raw_rows or not raw_rows[0]:
                continue

            headers = raw_rows[0]
            data_rows = raw_rows[1:]

            md_lines = ["| " + " | ".join(headers) + " |"]
            md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for r in data_rows:
                md_lines.append("| " + " | ".join(r) + " |")
            table_md = "\n".join(md_lines)

            elements.append(DocumentElement(
                element_id=f"{document_id}_docx_tbl{tbl_idx}",
                element_type=ElementType.TABLE,
                page_number=1,
                section_title=current_section,
                content_text=table_md,
                structured_data={
                    "headers": headers,
                    "rows": data_rows,
                    "row_count": len(data_rows),
                    "column_count": len(headers)
                }
            ))

        # Process inline images from docx package relationships
        img_counter = 0
        for rel in doc_obj.part.rels.values():
            if "image" in rel.target_ref:
                img_counter += 1
                img_bytes = rel.target_part.blob
                ext = rel.target_ref.split(".")[-1] if "." in rel.target_ref else "png"
                img_filename = f"{document_id}_docx_img{img_counter}.{ext}"
                img_path = os.path.join(settings.EXTRACTED_IMAGES_DIR, img_filename)

                with open(img_path, "wb") as f_img:
                    f_img.write(img_bytes)

                vision_result = vision_analyzer.analyze_image(
                    image_path=img_path,
                    context_caption=f"Image in DOCX {file_name}"
                )

                diag_text = (
                    f"[Diagram/Image: {vision_result.get('diagram_type', 'technical_diagram')}]\n"
                    f"Description: {vision_result.get('description', '')}\n"
                )
                if vision_result.get("components"):
                    diag_text += f"Components: {', '.join(vision_result['components'])}\n"
                if vision_result.get("connections"):
                    diag_text += f"Connections: {', '.join(vision_result['connections'])}\n"

                elements.append(DocumentElement(
                    element_id=f"{document_id}_docx_img{img_counter}",
                    element_type=ElementType.DIAGRAM if vision_result.get("diagram_type") != "general_image" else ElementType.IMAGE,
                    page_number=1,
                    section_title=current_section,
                    content_text=diag_text.strip(),
                    structured_data=vision_result,
                    image_path=img_path
                ))

        single_page = Page(
            page_number=1,
            elements=elements,
            page_text="\n".join([e.content_text for e in elements])
        )

        doc = Document(
            document_id=document_id,
            file_name=file_name,
            file_id=file_id,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            gdrive_url=gdrive_url,
            pages=[single_page],
            metadata={
                "pages_count": 1,
                "elements_count": len(elements)
            }
        )
        logger.info(f"Successfully parsed DOCX {file_name}: {len(elements)} elements.")
        return doc
