import os
import uuid
import pptx
from typing import List
from app.models.document import Document, Page, DocumentElement, ElementType
from app.parsing.base import BaseDocumentParser
from app.parsing.vision_analyzer import vision_analyzer
from app.core.config import settings
from app.core.logging import logger


class PPTXParser(BaseDocumentParser):
    """
    PPTX Document Parser extracting slides, titles, shapes, tables, and slide images.
    """

    def parse(self, file_path: str, file_id: str, file_name: str, gdrive_url: str = "") -> Document:
        logger.info(f"Parsing PPTX document: {file_name}")
        document_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, file_id or file_name))

        prs = pptx.Presentation(file_path)
        doc_pages: List[Page] = []
        total_elements = 0

        for slide_idx, slide in enumerate(prs.slides):
            page_num = slide_idx + 1
            page_elements: List[DocumentElement] = []
            slide_title = f"Slide {page_num}"

            # Check if slide has title
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()
                page_elements.append(DocumentElement(
                    element_id=f"{document_id}_s{page_num}_title",
                    element_type=ElementType.HEADING,
                    page_number=page_num,
                    section_title=slide_title,
                    content_text=slide_title
                ))
                total_elements += 1

            for shape_idx, shape in enumerate(slide.shapes):
                # Text frames
                if shape.has_text_frame and shape != slide.shapes.title:
                    stext = shape.text_frame.text.strip()
                    if stext:
                        page_elements.append(DocumentElement(
                            element_id=f"{document_id}_s{page_num}_sh{shape_idx}",
                            element_type=ElementType.TEXT,
                            page_number=page_num,
                            section_title=slide_title,
                            content_text=stext
                        ))
                        total_elements += 1

                # Tables
                elif shape.has_table:
                    table = shape.table
                    raw_rows = []
                    for row in table.rows:
                        raw_rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])

                    if raw_rows and raw_rows[0]:
                        headers = raw_rows[0]
                        data_rows = raw_rows[1:]

                        md_lines = ["| " + " | ".join(headers) + " |"]
                        md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                        for r in data_rows:
                            md_lines.append("| " + " | ".join(r) + " |")
                        table_md = "\n".join(md_lines)

                        page_elements.append(DocumentElement(
                            element_id=f"{document_id}_s{page_num}_tbl{shape_idx}",
                            element_type=ElementType.TABLE,
                            page_number=page_num,
                            section_title=slide_title,
                            content_text=table_md,
                            structured_data={
                                "headers": headers,
                                "rows": data_rows,
                                "row_count": len(data_rows),
                                "column_count": len(headers)
                            }
                        ))
                        total_elements += 1

                # Pictures
                elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        image_bytes = image.blob
                        ext = image.ext
                        img_filename = f"{document_id}_s{page_num}_img{shape_idx}.{ext}"
                        img_path = os.path.join(settings.EXTRACTED_IMAGES_DIR, img_filename)

                        with open(img_path, "wb") as f_img:
                            f_img.write(image_bytes)

                        vision_result = vision_analyzer.analyze_image(
                            image_path=img_path,
                            context_caption=f"Image on slide {page_num} ({slide_title}) in {file_name}"
                        )

                        diag_text = (
                            f"[Diagram/Image: {vision_result.get('diagram_type', 'technical_diagram')}]\n"
                            f"Description: {vision_result.get('description', '')}\n"
                        )
                        if vision_result.get("components"):
                            diag_text += f"Components: {', '.join(vision_result['components'])}\n"
                        if vision_result.get("connections"):
                            diag_text += f"Connections: {', '.join(vision_result['connections'])}\n"

                        page_elements.append(DocumentElement(
                            element_id=f"{document_id}_s{page_num}_img{shape_idx}",
                            element_type=ElementType.DIAGRAM if vision_result.get("diagram_type") != "general_image" else ElementType.IMAGE,
                            page_number=page_num,
                            section_title=slide_title,
                            content_text=diag_text.strip(),
                            structured_data=vision_result,
                            image_path=img_path
                        ))
                        total_elements += 1
                    except Exception as e:
                        logger.warning(f"Error processing picture shape on slide {page_num}: {e}")

            doc_pages.append(Page(
                page_number=page_num,
                elements=page_elements,
                page_text="\n".join([e.content_text for e in page_elements])
            ))

        doc = Document(
            document_id=document_id,
            file_name=file_name,
            file_id=file_id,
            mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            gdrive_url=gdrive_url,
            pages=doc_pages,
            metadata={
                "pages_count": len(doc_pages),
                "elements_count": total_elements
            }
        )
        logger.info(f"Successfully parsed PPTX {file_name}: {len(doc_pages)} slides, {total_elements} elements.")
        return doc
