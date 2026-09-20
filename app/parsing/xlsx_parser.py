import uuid
import openpyxl
from typing import List
from app.models.document import Document, Page, DocumentElement, ElementType
from app.parsing.base import BaseDocumentParser
from app.core.logging import logger


class XLSXParser(BaseDocumentParser):
    """
    XLSX Excel Document Parser preserving tabular structures, headers, and sheet sections.
    """

    def parse(self, file_path: str, file_id: str, file_name: str, gdrive_url: str = "") -> Document:
        logger.info(f"Parsing XLSX document: {file_name}")
        document_id = str(uuid.uuid5(uuid.NAMESPACE_DNS, file_id or file_name))

        wb = openpyxl.load_workbook(file_path, data_only=True)
        pages: List[Page] = []
        total_elements = 0

        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            page_num = sheet_idx + 1
            sheet = wb[sheet_name]
            page_elements: List[DocumentElement] = []

            # Add Sheet Heading
            page_elements.append(DocumentElement(
                element_id=f"{document_id}_p{page_num}_title",
                element_type=ElementType.HEADING,
                page_number=page_num,
                section_title=f"Sheet: {sheet_name}",
                content_text=f"Sheet: {sheet_name}"
            ))
            total_elements += 1

            rows = list(sheet.iter_rows(values_only=True))
            clean_rows = []
            for r in rows:
                if any(cell is not None and str(cell).strip() != "" for cell in r):
                    clean_rows.append([str(cell).strip() if cell is not None else "" for cell in r])

            if clean_rows:
                headers = clean_rows[0]
                data_rows = clean_rows[1:]

                md_lines = ["| " + " | ".join(headers) + " |"]
                md_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
                for r in data_rows:
                    md_lines.append("| " + " | ".join(r) + " |")
                table_md = "\n".join(md_lines)

                page_elements.append(DocumentElement(
                    element_id=f"{document_id}_p{page_num}_tbl",
                    element_type=ElementType.TABLE,
                    page_number=page_num,
                    section_title=f"Sheet: {sheet_name}",
                    content_text=table_md,
                    structured_data={
                        "sheet_name": sheet_name,
                        "headers": headers,
                        "rows": data_rows,
                        "row_count": len(data_rows),
                        "column_count": len(headers)
                    }
                ))
                total_elements += 1

            pages.append(Page(
                page_number=page_num,
                elements=page_elements,
                page_text="\n".join([e.content_text for e in page_elements])
            ))

        wb.close()

        doc = Document(
            document_id=document_id,
            file_name=file_name,
            file_id=file_id,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            gdrive_url=gdrive_url,
            pages=pages,
            metadata={
                "pages_count": len(pages),
                "elements_count": total_elements
            }
        )
        logger.info(f"Successfully parsed XLSX {file_name}: {len(pages)} sheets, {total_elements} elements.")
        return doc
